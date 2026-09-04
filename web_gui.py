import json, os, sys, threading, webbrowser, time, datetime, socket
from http.server import HTTPServer, ThreadingHTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pipeline import step_fetch, step_preprocess, step_llm_evaluate, step_postprocess, step_fetch_articles
from core import process_all, parse_input
from generate_queries import load_rows, generate_payload, generate_custom_queries

PORT = 18080
for port in range(18080, 18100):
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        s.bind(("127.0.0.1", port))
        s.close()
        PORT = port
        break
    except:
        continue

CHECKPOINT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "checkpoints")
os.makedirs(CHECKPOINT_DIR, exist_ok=True)

def save_checkpoint():
    """保存当前进度到检查点文件"""
    for k in range(1, 8):
        key = "step" + str(k) if k != 5 else "postprocess"
        data = state["results"].get(key)
        if data:
            try:
                cp = {"step": k, "data": data, "ts": datetime.datetime.now().isoformat()}
                with open(os.path.join(CHECKPOINT_DIR, f"step{k}.json"), "w", encoding="utf-8") as f:
                    json.dump(cp, f, ensure_ascii=False)
            except:
                pass
    # 保存非敏感配置（key 不落盘）
    try:
        public_cfg = {
            "base_url": state["config"].get("base_url", ""),
            "model": state["config"].get("model", ""),
            "provider": state["config"].get("provider", ""),
            "api_key": "",
            "tophub_key": "",
        }
        with open(os.path.join(CHECKPOINT_DIR, "config.json"), "w", encoding="utf-8") as f:
            json.dump(public_cfg, f, ensure_ascii=False)
    except:
        pass

def load_checkpoints():
    """从检查点文件恢复进度"""
    # 恢复配置
    cfg_path = os.path.join(CHECKPOINT_DIR, "config.json")
    if os.path.exists(cfg_path):
        try:
            with open(cfg_path, "r", encoding="utf-8") as f:
                saved = json.load(f)
            state["config"]["base_url"] = saved.get("base_url", "")
            state["config"]["model"] = saved.get("model", "")
            state["config"]["provider"] = saved.get("provider", "")
            state["config"]["api_key"] = ""
            state["config"]["tophub_key"] = ""
        except:
            pass
    # 恢复步骤结果
    for k in range(1, 8):
        cp_path = os.path.join(CHECKPOINT_DIR, f"step{k}.json")
        if os.path.exists(cp_path):
            try:
                with open(cp_path, "r", encoding="utf-8") as f:
                    cp = json.load(f)
                key = "step" + str(k) if k != 5 else "postprocess"
                state["results"][key] = cp["data"]
                state["steps"][k] = "done"
                add_log(f"检测到检查点: Step {k} 已完成", "info")
            except:
                os.remove(cp_path)  # 损坏文件直接删除

state = {
    "config": {"tophub_key": "", "api_key": "", "base_url": "", "model": "", "provider": ""},
    "steps": {1: "idle", 2: "idle", 3: "idle", 4: "idle", 5: "idle", 6: "idle", 7: "idle"},
    "step7_config": {"total": 10, "black": 6, "white": 2, "gray": 2, "rounds": "单轮", "multi_rounds": 2},
    "results": {},
    "logs": [],
    "running": False,
}

def add_log(msg, level="info"):
    ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    state["logs"].append({"time": ts, "msg": msg, "level": level})
    if len(state["logs"]) > 1000:
        state["logs"] = state["logs"][-500:]

def reset_pipeline():
    for k in range(1, 8):
        state["steps"][k] = "idle"
    state["results"] = {}
    # 清除检查点文件
    for k in range(1, 8):
        cp = os.path.join(CHECKPOINT_DIR, f"step{k}.json")
        if os.path.exists(cp):
            try: os.remove(cp)
            except: pass
    add_log("Pipeline reset", "info")

def _merge_full_text():
    """把 Step 6 抓到的原文合入后处理结果；未跑 Step 6 时退回 Step 3 的原文"""
    rows = []
    if state["results"].get("postprocess"):
        rows = json.loads(state["results"]["postprocess"]["result_json"])
    text_map = {}
    step6 = state["results"].get("step6")
    if isinstance(step6, dict):
        for item in step6.get("enriched", []):
            try:
                idx = int(item.get("idx"))
            except Exception:
                continue
            text_map[idx] = {
                "full_text": item.get("full_text"),
                "fetch_status": item.get("fetch_status", ""),
                "matched_url": item.get("matched_url", ""),
                "verdict": item.get("最终判定"),
                "reason": item.get("原文回填", ""),
            }
    step3 = state["results"].get("step3")
    if isinstance(step3, dict):
        for item in step3.get("items", []):
            try:
                idx = int(item.get("idx"))
            except Exception:
                continue
            if idx not in text_map:
                text_map[idx] = {
                    "full_text": item.get("full_text"),
                    "fetch_status": item.get("fetch_status", ""),
                    "matched_url": item.get("matched_url", ""),
                    "verdict": None,
                    "reason": "",
                }
    for row in rows:
        try:
            row_idx = int(row.get("idx") or -1)
        except Exception:
            row_idx = -1
        info = text_map.get(row_idx, {})
        row["full_text"] = info.get("full_text")
        row["fetch_status"] = info.get("fetch_status", "")
        row["matched_url"] = info.get("matched_url", "")
        if info.get("verdict"):
            row["最终判定"] = info.get("verdict")
        if info.get("reason"):
            row["原文回填"] = info.get("reason")
    return rows

def _make_report_workbook():
    """生成全量条目的 Excel 报告，原文全文直接放在主表，保存到 outputs 并返回路径"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    import datetime

    items = _merge_full_text()
    wb = Workbook()
    ws = wb.active
    ws.title = "报告"

    hf = Font(bold=True, size=11, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="2F5496")
    gfill = PatternFill("solid", fgColor="E2EFDA")
    wfill = PatternFill("solid", fgColor="FFF2CC")
    bfill = PatternFill("solid", fgColor="FCE4EC")
    tb = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

    headers = ["idx", "来源", "标题", "评分", "最终判定", "分类", "冲突", "极性", "主体", "质量", "入选理由", "URL", "原文链接", "原文状态", "原文全文"]
    ncol = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws["A1"].value = "舆情清洗报告"
    ws["A1"].font = Font(bold=True, size=14, color="2F5496")
    ws["A1"].alignment = Alignment(horizontal="center")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = hf; cell.fill = hfill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = tb

    for col, width in enumerate([6, 14, 40, 8, 10, 8, 8, 8, 14, 12, 40, 50, 50, 20, 80], 1):
        ws.column_dimensions[chr(64 + col)].width = width

    for ri, item in enumerate(items, 4):
        q = item.get("fetch_quality", "")
        if q == "real_article": fill = gfill
        elif q == "partial": fill = wfill
        else: fill = bfill
        vals = [
            item.get("idx"), item.get("来源"), item.get("标题"),
            item.get("综合评分"), item.get("最终判定"), item.get("分类"),
            item.get("冲突强度"), item.get("极性"),
            item.get("主体", ""), q,
            (item.get("入选理由") or "")[:40], item.get("url", ""),
            item.get("matched_url") or "", item.get("fetch_status", ""),
            item.get("full_text") or ("未获取: " + (item.get("fetch_status") or "未知原因")),
        ]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=col, value=val)
            cell.border = tb
            cell.alignment = Alignment(vertical="top", wrap_text=(col >= 11))
            if col == 10: cell.fill = fill

    sr = len(items) + 5
    ws.merge_cells(start_row=sr, start_column=1, end_row=sr, end_column=ncol)
    fetched = len([i for i in items if i.get("full_text")])
    sc = len([i for i in items if i.get("最终判定") == "入选"])
    wc = len([i for i in items if i.get("最终判定") == "待观察"])
    rc = len(items) - sc - wc
    ws[f"A{sr}"].value = f"入选: {sc}条 | 待观察: {wc}条 | 淘汰: {rc}条 | 合计: {len(items)}条 | 已获取原文: {fetched}条"
    ws[f"A{sr}"].font = Font(bold=True, size=10, color="2F5496")

    ts = datetime.datetime.now().strftime("%Y%m%d%H%M")
    out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "outputs")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f"pipeline_result_{ts}.xlsx")
    wb.save(out_path)
    return out_path

def _parse_category_rules():
    """从 clean_prompt.txt 4a 抽取 A-L 品类定义，返回结构化列表"""
    import re
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "clean_prompt.txt")
    with open(path, "r", encoding="utf-8") as f:
        text = f.read()
    a = text.find("4a.")
    b = text.find("4b.")
    sec = text[a:b] if 0 <= a < b else ""
    categories = []
    cur = None
    for ln in sec.splitlines():
        m = re.match(r"\s*([A-L])-([^\s：:]+)[：:]?\s*(.*)", ln)
        if m:
            if cur:
                categories.append(cur)
            name = m.group(2).strip("（）()")
            items = [m.group(3).strip()] if m.group(3).strip() else []
            cur = {"letter": m.group(1), "name": name, "items": items}
        elif cur and ln.strip():
            cur["items"].append(ln.strip())
    if cur:
        categories.append(cur)
    return categories

def _parse_changelog():
    """解析 docs/CHANGELOG.md，按条目返回（文件内已倒序）"""
    import re
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "docs", "CHANGELOG.md")
    if not os.path.exists(path):
        return []
    with open(path, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()
    entries = []
    cur = None
    for ln in lines:
        m = re.match(r"^##\s+(\d{4}-\d{2}-\d{2})\s+(.*)$", ln.strip())
        if m:
            if cur:
                entries.append(cur)
            cur = {"date": m.group(1), "title": m.group(2).strip(), "items": []}
        elif cur and ln.strip().startswith("- "):
            cur["items"].append(ln.strip()[2:].strip())
    if cur:
        entries.append(cur)
    return entries

class Handler(BaseHTTPRequestHandler):
    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path
        if path == "/":
            self._serve_file("index.html")
        elif path == "/api/status":
            self._json({
                "config": {
                    "base_url": state["config"].get("base_url", ""),
                    "model": state["config"].get("model", ""),
                    "provider": state["config"].get("provider", ""),
                    "api_key_set": bool(state["config"].get("api_key")),
                    "tophub_key_set": bool(state["config"].get("tophub_key")),
                },
                "steps": state["steps"],
                "running": state["running"],
                "logs": state["logs"][-100:],
                "results": {k: bool(v) for k, v in state["results"].items()},
            })
        elif path == "/api/export":
            if not state["results"].get("postprocess"):
                self._json({"error": "No results to export"}); return
            payload = dict(state["results"]["postprocess"])
            payload["result_json"] = json.dumps(_merge_full_text(), ensure_ascii=False)
            self._json(payload)
        elif path == "/api/export/xlsx":
            if not state["results"].get("postprocess"):
                self._json({"error": "No results to export"}); return
            self._serve_xlsx()
        elif path == "/api/export/queries":
            if not state["results"].get("step7"):
                self._json({"error": "No query results"}); return
            data = json.dumps(state["results"]["step7"], ensure_ascii=False).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Disposition", 'attachment; filename="query_results.json"')
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
        elif path == "/api/export/queries/xlsx":
            if not state["results"].get("step7"):
                self._json({"error": "No query results"}); return
            try:
                from generate_queries import save_queries_xlsx
                out_path = save_queries_xlsx(state["results"]["step7"])
                with open(out_path, "rb") as f:
                    file_data = f.read()
                fname = os.path.basename(out_path)
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", 'attachment; filename="%s"' % fname)
                self.send_header("Content-Length", str(len(file_data)))
                self.end_headers()
                self.wfile.write(file_data)
            except Exception as e:
                self._json({"error": "Excel export failed: %s" % str(e)[:200]})
        elif path == "/api/export/custom/xlsx":
            batch = state.get("custom_batch") or {}
            if not batch.get("items"):
                self._json({"error": "No custom query results"}); return
            try:
                from generate_queries import save_custom_xlsx
                out_path = save_custom_xlsx(batch["items"])
                with open(out_path, "rb") as f:
                    file_data = f.read()
                fname = os.path.basename(out_path)
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", 'attachment; filename="%s"' % fname)
                self.send_header("Content-Length", str(len(file_data)))
                self.end_headers()
                self.wfile.write(file_data)
            except Exception as e:
                self._json({"error": "Excel export failed: %s" % str(e)[:200]})
        elif path == "/api/rules":
            try:
                categories = _parse_category_rules()
            except Exception as e:
                categories = []
                self._json({"error": "读取规则失败: %s" % str(e)[:200]}); return
            self._json({"categories": categories})
        elif path == "/api/changelog":
            try:
                changelog = _parse_changelog()
            except Exception as e:
                changelog = []
                self._json({"error": "读取更新日志失败: %s" % str(e)[:200]}); return
            self._json({"entries": changelog})
        elif path == "/api/step7/config":
            items = []
            try:
                from generate_queries import load_rows
                rows = [r for r in load_rows() if r.get("最终判定") == "入选"]
                items = [{"idx": r.get("idx"), "来源": r.get("来源"), "标题": r.get("标题"), "分类": r.get("分类")} for r in rows]
            except Exception:
                items = []
            self._json({"config": state["step7_config"], "items": items})
        elif path.startswith("/api/step/"):
            parts = path.split("/")
            if len(parts) >= 4 and parts[3].isdigit():
                step_num = int(parts[3])
                data = state["results"].get("step" + str(step_num))
                if data: self._json(data)
                else: self._json({"error": "No data"})
            else: self.send_error(404)
        else: self.send_error(404)

    def do_POST(self):
        parsed = urlparse(self.path)
        path = parsed.path
        length = int(self.headers.get("Content-Length", 0))
        body = self.rfile.read(length).decode("utf-8") if length else "{}"
        try: data = json.loads(body)
        except: data = {}

        if path == "/api/configure":
            if data.get("tophub_key"):
                state["config"]["tophub_key"] = data["tophub_key"].strip()
            if data.get("api_key"):
                state["config"]["api_key"] = data["api_key"].strip()
            if "base_url" in data: state["config"]["base_url"] = data["base_url"].strip()
            if "model" in data: state["config"]["model"] = data["model"].strip()
            if "provider" in data: state["config"]["provider"] = data["provider"].strip()
            from llm_client import set_llm_config
            set_llm_config(
                base_url=state["config"]["base_url"],
                api_key=state["config"]["api_key"],
                model=state["config"]["model"],
            )
            add_log("Configuration updated (内存生效，不落盘)", "info")
            reset_pipeline()
            self._json({"ok": True, "api_key_set": bool(state["config"]["api_key"]), "tophub_key_set": bool(state["config"]["tophub_key"])})

        elif path == "/api/run/step/1":
            if state["running"]: self._json({"error": "Already running"}); return
            if not state["config"]["tophub_key"]: self._json({"error": "Tophub Key not configured"}); return
            self._json({"ok": True, "msg": "Step 1 started"})
            threading.Thread(target=_run_step1, daemon=True).start()

        elif path == "/api/run/step/2":
            if state["running"]: self._json({"error": "Already running"}); return
            if not state["results"].get("step1"): self._json({"error": "Step 1 not completed"}); return
            self._json({"ok": True, "msg": "Step 2 started"})
            threading.Thread(target=_run_step2, daemon=True).start()

        elif path == "/api/run/step/3":
            if state["running"]: self._json({"error": "Already running"}); return
            if not state["results"].get("step2"): self._json({"error": "Step 2 not completed"}); return
            self._json({"ok": True, "msg": "Step 3 started"})
            threading.Thread(target=_run_step3, daemon=True).start()

        elif path == "/api/run/step/4":
            if state["running"]: self._json({"error": "Already running"}); return
            if not state["results"].get("step3"): self._json({"error": "Step 3 not completed"}); return
            self._json({"ok": True, "msg": "Step 4 started"})
            threading.Thread(target=_run_step4, daemon=True).start()

        elif path == "/api/run/step/5":
            if state["running"]: self._json({"error": "Already running"}); return
            if not state["results"].get("step4"): self._json({"error": "Step 4 not completed"}); return
            self._json({"ok": True, "msg": "Step 5 started"})
            threading.Thread(target=_run_step5, daemon=True).start()

        elif path == "/api/run/step/6":
            if state["running"]: self._json({"error": "Already running"}); return
            if not state["results"].get("postprocess"): self._json({"error": "Step 5 not completed"}); return
            self._json({"ok": True, "msg": "Step 6 started"})
            threading.Thread(target=_run_step6, daemon=True).start()

        elif path == "/api/run/step/7":
            if state["running"]: self._json({"error": "Already running"}); return
            if not state["results"].get("step6"): self._json({"error": "Step 6 not completed"}); return
            self._json({"ok": True, "msg": "Step 7 started"})
            threading.Thread(target=_run_step7, daemon=True).start()

        elif path == "/api/step7/config":
            try:
                total = int(data.get("total") or 10)
                black = int(data.get("black") or 0)
                white = int(data.get("white") or 0)
                gray = int(data.get("gray") or 0)
                multi_rounds = int(data.get("multi_rounds") or 2)
            except Exception:
                self._json({"error": "数量参数必须是整数"}); return
            rounds = str(data.get("rounds") or "单轮+多轮")
            if rounds not in ("单轮", "多轮", "单轮+多轮"):
                self._json({"error": "轮次类型不合法"}); return
            if multi_rounds < 2 or multi_rounds > 10:
                self._json({"error": "多轮轮次必须在 2-10 之间"}); return
            if black + white + gray != total:
                self._json({"error": "黑/白/灰数量之和必须等于总数"}); return
            state["step7_config"] = {"total": total, "black": black, "white": white, "gray": gray, "rounds": rounds, "multi_rounds": multi_rounds}
            self._json({"ok": True, "config": state["step7_config"]})

        elif path == "/api/step7/run":
            if state["running"]: self._json({"error": "Already running"}); return
            if not state["results"].get("step6"): self._json({"error": "Step 6 not completed"}); return
            cfg = state["step7_config"]
            if int(cfg["black"]) + int(cfg["white"]) + int(cfg["gray"]) != int(cfg["total"]):
                self._json({"error": "Step 7 配置无效，请先保存配置"}); return
            self._json({"ok": True, "msg": "Step 7 started"})
            threading.Thread(target=_run_step7, daemon=True).start()

        elif path == "/api/analyze/custom":
            if state["running"]:
                self._json({"error": "Already running"}); return
            title = (data.get("title") or "").strip()
            content = (data.get("content") or "").strip()
            if not title or not content:
                self._json({"error": "标题和内容不能为空"}); return
            try:
                total = int(data.get("total") or 10)
                black = int(data.get("black") or 0)
                white = int(data.get("white") or 0)
                gray = int(data.get("gray") or 0)
                multi_rounds = int(data.get("multi_rounds") or 2)
            except Exception:
                self._json({"error": "数量参数必须是整数"}); return
            rounds = str(data.get("rounds") or "单轮")
            if rounds not in ("单轮", "多轮", "单轮+多轮"):
                self._json({"error": "轮次类型不合法"}); return
            if multi_rounds < 2 or multi_rounds > 10:
                self._json({"error": "多轮轮次必须在 2-10 之间"}); return
            if black + white + gray != total:
                self._json({"error": "黑/白/灰数量之和必须等于 Query 总数"}); return
            result = generate_custom_queries(
                title,
                content,
                source=(data.get("source") or "").strip(),
                category_override=(data.get("category") or "").strip(),
                total=total,
                black=black,
                white=white,
                gray=gray,
                extra=(data.get("extra") or "").strip(),
                rounds=rounds,
                multi_rounds=multi_rounds,
            )
            self._json(result)

        elif path == "/api/analyze/custom/batch":
            if state["running"]:
                self._json({"error": "Already running"}); return
            items = data.get("items") or []
            if not isinstance(items, list) or not items:
                self._json({"error": "至少需要一条舆情"}); return
            if len(items) > 50:
                self._json({"error": "单次最多 50 条舆情"}); return
            try:
                total = int(data.get("total") or 10)
                black = int(data.get("black") or 0)
                white = int(data.get("white") or 0)
                gray = int(data.get("gray") or 0)
                multi_rounds = int(data.get("multi_rounds") or 2)
            except Exception:
                self._json({"error": "数量参数必须是整数"}); return
            rounds = str(data.get("rounds") or "单轮")
            if rounds not in ("单轮", "多轮", "单轮+多轮"):
                self._json({"error": "轮次类型不合法"}); return
            if multi_rounds < 2 or multi_rounds > 10:
                self._json({"error": "多轮轮次必须在 2-10 之间"}); return
            if black + white + gray != total:
                self._json({"error": "黑/白/灰数量之和必须等于 Query 总数"}); return
            cfg = {"total": total, "black": black, "white": white, "gray": gray,
                   "rounds": rounds, "multi_rounds": multi_rounds, "extra": (data.get("extra") or "").strip()}
            from generate_queries import generate_custom_batch
            payload = generate_custom_batch(items, cfg, lambda m: add_log("自定义出题: " + m, "info"))
            payload["config"] = cfg
            state["custom_batch"] = payload
            self._json(payload)

        elif path == "/api/analyze/custom/classify":
            if state["running"]:
                self._json({"error": "Already running"}); return
            items = data.get("items") or []
            if not isinstance(items, list) or not items:
                self._json({"error": "至少需要一条舆情"}); return
            if len(items) > 50:
                self._json({"error": "单次最多 50 条舆情"}); return
            from generate_queries import classify_custom_items
            payload = classify_custom_items(items, lambda m: add_log("自定义打标: " + m, "info"))
            state["custom_classified"] = payload
            self._json(payload)

        elif path == "/api/run/all":
            if state["running"]: self._json({"error": "Already running"}); return
            self._json({"ok": True, "msg": "Full pipeline started"})
            threading.Thread(target=_run_all, daemon=True).start()


        elif path == "/api/export/xlsx":
            if not state["results"].get("postprocess"):
                self._json({"error": "No results to export"})
                return
            try:
                out_path = _make_report_workbook()
                with open(out_path, "rb") as f:
                    file_data = f.read()
                fname = os.path.basename(out_path)
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f'attachment; filename="{fname}"')
                self.send_header("Content-Length", str(len(file_data)))
                self.end_headers()
                self.wfile.write(file_data)
            except Exception as e:
                self._json({"error": f"Excel export failed: {str(e)[:200]}"})

        elif path == "/api/reset":
            reset_pipeline()
            self._json({"ok": True})
        elif path == "/api/checkpoint":
            # 返回检查点状态
            cps = {}
            for k in range(1, 8):
                cp = os.path.join(CHECKPOINT_DIR, f"step{k}.json")
                cps[f"step{k}"] = os.path.exists(cp)
            cps["config"] = os.path.exists(os.path.join(CHECKPOINT_DIR, "config.json"))
            self._json(cps)
        else: self.send_error(404)

    def _serve_file(self, filename):
        path = os.path.join(os.path.dirname(__file__), filename)
        if not os.path.exists(path): self.send_error(404); return
        with open(path, "rb") as f: data = f.read()
        self.send_response(200)
        if filename.endswith(".html"): self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _json(self, obj):
        data = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _serve_xlsx(self):
        try:
            out_path = _make_report_workbook()
            with open(out_path, "rb") as f:
                file_data = f.read()
            fname = os.path.basename(out_path)
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", f'attachment; filename="{fname}"')
            self.send_header("Content-Length", str(len(file_data)))
            self.end_headers()
            self.wfile.write(file_data)
        except Exception as e:
            self._json({"error": f"Excel export failed: {str(e)[:200]}"})

    def log_message(self, format, *args): pass

def _run_step1():
    state["running"] = True; state["steps"][1] = "running"
    add_log("Step 1: Starting fetch from Tophub...", "info")
    try:
        result = step_fetch(state["config"]["tophub_key"], lambda m: add_log("  " + m, "info"))
        state["results"]["step1"] = result; state["steps"][1] = "done"
        save_checkpoint()
        add_log("Step 1: Complete", "success")
    except Exception as e: state["steps"][1] = "error"; add_log("Step 1: Error - " + str(e), "error")
    finally: state["running"] = False

def _run_step2():
    state["running"] = True; state["steps"][2] = "running"
    add_log("Step 2: Starting preprocessing...", "info")
    try:
        result = step_preprocess(state["results"]["step1"], lambda m: add_log("  " + m, "info"))
        state["results"]["step2"] = result; state["steps"][2] = "done"
        save_checkpoint()
        add_log("Step 2: Complete", "success")
    except Exception as e: state["steps"][2] = "error"; add_log("Step 2: Error - " + str(e), "error")
    finally: state["running"] = False

def _run_step3():
    state["running"] = True; state["steps"][3] = "running"
    add_log("Step 3: Starting article fetch...", "info")
    try:
        items = json.loads(state["results"]["step2"]["items_json"])
        result = step_fetch_articles(items, lambda m: add_log("  " + m, "info"))
        state["results"]["step3"] = result; state["steps"][3] = "done"
        save_checkpoint()
        add_log("Step 3: Complete - " + str(result["stats"]["success"]) + "/" + str(result["stats"]["total"]) + " fetched", "success")
    except Exception as e: state["steps"][3] = "error"; add_log("Step 3: Error - " + str(e), "error")
    finally: state["running"] = False

def _run_step4():
    state["running"] = True; state["steps"][4] = "running"
    add_log("Step 4: Starting LLM evaluation...", "info")
    try:
        fetch_results = state["results"].get("step3")
        result = step_llm_evaluate(state["results"]["step2"]["to_llm_text"], lambda m: add_log("  " + m, "info"), fetch_results=fetch_results)
        state["results"]["step4"] = result
        if result.get("error"):
            add_log("Step 4: " + result["error"], "error")
            state["steps"][4] = "error"
        else:
            state["steps"][4] = "done"
            save_checkpoint()
            add_log("Step 4: Complete", "success")
    except Exception as e: state["steps"][4] = "error"; add_log("Step 4: Error - " + str(e), "error")
    finally: state["running"] = False

def _run_step5():
    state["running"] = True; state["steps"][5] = "running"
    add_log("Step 5: Starting postprocessing...", "info")
    try:
        fetch_items = state["results"].get("step3")
        result = step_postprocess(
            state["results"]["step4"]["raw_output"],
            state["results"]["step2"]["items_json"],
            state["results"]["step2"]["auto_out_json"],
            lambda m: add_log("  " + m, "info"),
            fetch_items=fetch_items)
        state["results"]["step5"] = result; state["results"]["postprocess"] = result
        state["steps"][5] = "done"
        save_checkpoint()
        r = json.loads(result["result_json"])
        selected = len([i for i in r if i.get("最终判定") == "入选"])
        watching = len([i for i in r if i.get("\u6700\u7ec8\u5224\u5b9a") == "待观察"])
        add_log("Step 5: Complete - Selected: " + str(selected) + ", Watching: " + str(watching), "success")
    except Exception as e: state["steps"][5] = "error"; add_log("Step 5: Error - " + str(e), "error")
    finally: state["running"] = False

def _run_step6():
    state["running"] = True; state["steps"][6] = "running"
    add_log("Step 6: Starting article enrichment...", "info")
    try:
        from pipeline import step_enrich_articles
        result = step_enrich_articles(state["results"]["postprocess"], lambda m: add_log("  " + m, "info"), state["results"].get("step3"))
        state["results"]["step6"] = result; state["steps"][6] = "done"
        save_checkpoint()
        s = result["stats"]
        add_log(f"Step 6: Complete - {s['success']}/{s['total']} enriched", "success")
    except Exception as e: state["steps"][6] = "error"; add_log("Step 6: Error - " + str(e), "error")
    finally: state["running"] = False

def _run_step7():
    state["running"] = True; state["steps"][7] = "running"
    add_log("Step 7: Starting query generation...", "info")
    try:
        from generate_queries import load_rows, generate_payload
        rows = [r for r in load_rows() if r.get("最终判定") == "入选"]
        rows.sort(key=lambda x: int(x.get("idx") or 0))
        payload = generate_payload(rows, state.get("step7_config"), lambda m: add_log("  " + m, "info"))
        try:
            from generate_queries import save_queries_xlsx
            payload["xlsx"] = save_queries_xlsx(payload)
        except Exception as e:
            add_log("Step 7: Excel 生成失败 - " + str(e), "error")
        state["results"]["step7"] = payload
        state["steps"][7] = "done"
        save_checkpoint()
        add_log("Step 7: Complete - " + json.dumps(payload["stats"], ensure_ascii=False), "success")
    except Exception as e: state["steps"][7] = "error"; add_log("Step 7: Error - " + str(e), "error")
    finally: state["running"] = False

def _run_all():
    """从第一个未完成的步骤开始连续运行，支持断点续转"""
    if state["steps"][1] != "done": _run_step1()
    if state["steps"][1] == "done" and state["steps"][2] != "done": _run_step2()
    if state["steps"][2] == "done" and state["steps"][3] != "done": _run_step3()
    if state["steps"][3] == "done" and state["steps"][4] != "done": _run_step4()
    if state["steps"][4] == "done" and state["steps"][5] != "done": _run_step5()
    if state["steps"][5] == "done" and state["steps"][6] != "done": _run_step6()
    if state["steps"][6] == "done" and state["steps"][7] != "done":
        add_log("Step 6 完成，Step 7 等待人工配置后运行", "info")
    if state["steps"][7] == "done":
        add_log("Full pipeline complete", "success")

def main():
    # 加载检查点
    load_checkpoints()
    ThreadingHTTPServer.allow_reuse_address = True
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("10.255.255.255", 1))
        lan_ip = s.getsockname()[0]
        s.close()
    except Exception:
        lan_ip = "127.0.0.1"
    server = ThreadingHTTPServer(("0.0.0.0", PORT), Handler)
    print()
    print("=== Server started ===")
    print("本机访问: http://127.0.0.1:" + str(PORT))
    print("内网访问: http://" + str(lan_ip) + ":" + str(PORT))
    print("Press Ctrl+C to stop")
    print()
    add_log("Server started on http://" + str(lan_ip) + ":" + str(PORT), "info")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nServer stopped")

if __name__ == "__main__":
    main()
